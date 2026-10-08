from datetime import datetime, time, timedelta
from io import BytesIO
import json
from pathlib import Path
from unittest.mock import Mock
from xml.etree import ElementTree
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
import pandas as pd
import pytest

from backend.app.config import Settings
from backend.app.datasets import MAX_UPLOAD_BYTES, parse_csv
from backend.app.excel import ENGINE_OPTIONS, MAX_UNCOMPRESSED_BYTES, UNSUPPORTED_EXCEL, parse_xlsx
from backend.app.main import create_app


MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def workbook_bytes(sheets):
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets:
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def excel_client(llm=None):
    settings = Settings(
        gemini_api_key="fixture", gemini_backup_api_key="",
        csv_path=Path(__file__).parents[1] / "data" / "sales_data.csv",
        rate_limit_per_minute=100,
    )
    return TestClient(create_app(settings, llm_client=llm or Mock()))


def upload(client, content, filename="customers.xlsx"):
    return client.post("/api/datasets", params={"filename": filename}, content=content,
                       headers={"Content-Type": MIME})


def test_single_sheet_has_csv_shape_schema_rows_and_query():
    llm = Mock()
    llm.generate.return_value = json.dumps({"aggregations": [{"column": "spend", "func": "sum", "alias": "total"}]})
    content = workbook_bytes([("Customers", [
        ["customer", "joined_date", "spend"],
        ["Ada", datetime(2026, 1, 2), 12.5], ["Lin", datetime(2026, 2, 3), 20.0],
    ])])
    with excel_client(llm) as client:
        csv = client.post("/api/datasets?filename=customers.csv",
                          content=b"customer,joined_date,spend\nAda,2026-01-02,12.5\nLin,2026-02-03,20.0\n").json()
        response = upload(client, content)
        assert response.status_code == 201
        excel = response.json()
        assert excel["name"] == "customers.xlsx – Customers"
        assert len(excel["dataset_id"]) == 32
        assert set(excel) == set(csv)
        for key in ("columns", "column_names", "rows", "row_count"):
            assert excel[key] == csv[key]
        assert client.get("/api/schema", params={"dataset_id": excel["dataset_id"]}).json() == client.get("/api/schema", params={"dataset_id": csv["dataset_id"]}).json()
        rows = client.get(f'/api/datasets/{excel["dataset_id"]}/rows', params={"offset": 1, "limit": 1}).json()
        assert rows["rows"] == csv["rows"][1:]
        assert rows["row_count"] == 2
        assert rows["has_more"] is False
        csv_query = client.post("/api/query", json={"question": "total spend", "dataset_id": csv["dataset_id"]})
        excel_query = client.post("/api/query", json={"question": "total spend", "dataset_id": excel["dataset_id"]})
        assert csv_query.status_code == excel_query.status_code == 200
        assert excel_query.json() == csv_query.json()
        assert excel_query.json()["rows"] == [{"total": 32.5}]
        assert llm.generate.call_count == 2


def test_multiple_sheets_have_independent_datasets_and_skip_empty_sheets():
    content = workbook_bytes([
        ("First", [["value"], [1]]), ("Empty", [[], [None]]),
        ("Second", [["value"], [2]]), ("Third", [["value"], [3]]),
    ])
    with excel_client() as client:
        response = upload(client, content)
        assert response.status_code == 201
        body = response.json()
        assert [sheet["name"] for sheet in body["datasets"]] == [f"customers.xlsx – {name}" for name in ("First", "Second", "Third")]
        assert len({sheet["dataset_id"] for sheet in body["datasets"]}) == 3
        assert body["dataset_id"] == body["datasets"][0]["dataset_id"]
        assert body["rows"] == [{"value": 1}]
        for value, dataset in enumerate(body["datasets"], 1):
            assert set(dataset) == {"dataset_id", "name"}
            assert client.get(f'/api/datasets/{dataset["dataset_id"]}/rows').json()["rows"] == [{"value": value}]
        assert client.app.state.datasets.get("default").name == "Bundled sales data"


def test_first_nonempty_row_is_header_and_empty_rows_columns_are_dropped():
    content = workbook_bytes([("Padded", [
        [], [None, "  ", None], [None, " name ", None, "value", None],
        [None, "Ada", None, 3, None], [], [None, "Lin", None, 4, None], [],
    ])])
    with excel_client() as client:
        response = upload(client, content, "PADDED.XLSX")
    assert response.status_code == 201
    assert response.json()["column_names"] == ["name", "value"]
    assert response.json()["rows"] == [{"name": "Ada", "value": 3}, {"name": "Lin", "value": 4}]
    assert response.json()["row_count"] == 2


@pytest.mark.parametrize("header", [
    ["name", "name"], ["name", " name "], [None, "value"],
    ["Unnamed: 0", "value"], [1, 2], [" ", "value"],
], ids=["duplicate", "trimmed-duplicate", "missing", "unnamed", "numeric", "blank"])
def test_bad_headers_reject_entire_workbook_without_storing_sheets(header):
    content = workbook_bytes([("Good", [["value"], [1]]), ("Bad", [header, ["Ada", 3]])])
    with excel_client() as client:
        response = upload(client, content)
        assert len(client.app.state.datasets._datasets) == 1
    assert response.status_code == 400
    assert response.json() == {"error": 'Sheet "Bad" must have a header row with non-empty, unique column names.'}


@pytest.mark.parametrize("filename", ["old.xls", "macros.xlsm", "OLD.XLS", "MACROS.XLSM"])
def test_unsupported_extensions(filename):
    with excel_client() as client:
        response = upload(client, b"unused", filename)
    assert response.status_code == 400
    assert response.json() == {"error": UNSUPPORTED_EXCEL}


@pytest.mark.parametrize("content,message", [
    (b"", "The uploaded workbook is empty."),
    (b"not a zip", "The file could not be parsed as an .xlsx workbook."),
], ids=["empty", "invalid-zip"])
def test_invalid_workbooks(content, message):
    with excel_client() as client:
        response = upload(client, content)
    assert response.status_code == 400
    assert response.json() == {"error": message}


def test_no_nonempty_sheets():
    with excel_client() as client:
        response = upload(client, workbook_bytes([("Empty", [[]])]))
    assert response.status_code == 400
    assert response.json() == {"error": "The workbook has no non-empty sheets."}


def test_ten_sheet_limit_includes_empty_sheets():
    with excel_client() as client:
        accepted = upload(client, workbook_bytes([(f"Sheet{i}", [["value"], [i]]) for i in range(10)]))
        assert accepted.status_code == 201
        assert len(accepted.json()["datasets"]) == 10
        response = upload(client, workbook_bytes([(f"Sheet{i}", []) for i in range(11)]))
    assert response.status_code == 400
    assert response.json() == {"error": "The workbook exceeds the 10-sheet limit."}


def test_upload_size_limit():
    with excel_client() as client:
        response = upload(client, b"x" * (MAX_UPLOAD_BYTES + 1))
    assert response.status_code == 413
    assert response.json() == {"error": "The workbook exceeds the 10 MB upload limit."}


def test_zip_bomb_is_rejected_before_excel_parser(monkeypatch):
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        with archive.open("xl/worksheets/sheet1.xml", "w") as entry:
            for _ in range(MAX_UNCOMPRESSED_BYTES // (1024 * 1024) + 1):
                entry.write(b"0" * (1024 * 1024))
    content = output.getvalue()
    assert len(content) < MAX_UPLOAD_BYTES
    parser = Mock(side_effect=AssertionError("Excel parser must not run"))
    monkeypatch.setattr("backend.app.excel.pd.ExcelFile", parser)
    with excel_client() as client:
        response = upload(client, content)
    assert response.status_code == 400
    assert response.json() == {"error": "The workbook exceeds the 100 MB uncompressed size limit."}
    parser.assert_not_called()


def test_column_limit():
    content = workbook_bytes([("Wide", [[f"c{i}" for i in range(101)], list(range(101))])])
    with excel_client() as client:
        response = upload(client, content)
    assert response.status_code == 400
    assert response.json() == {"error": 'Sheet "Wide" exceeds the 100-column limit.'}


def test_row_limit():
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Long")
    sheet.append(["value"])
    for _ in range(100001):
        sheet.append([1])
    output = BytesIO()
    workbook.save(output)
    with excel_client() as client:
        response = upload(client, output.getvalue())
    assert response.status_code == 400
    assert response.json() == {"error": 'Sheet "Long" exceeds the 100,000-row limit.'}


def test_sparse_cells_rejected_before_allocating_rectangular_frame(monkeypatch):
    workbook = Workbook()
    workbook.active["A1"] = "value"
    workbook.active["XFD1048576"] = 1
    output = BytesIO()
    workbook.save(output)
    parser = Mock(side_effect=AssertionError("Excel parser must not run"))
    monkeypatch.setattr("backend.app.excel.pd.ExcelFile", parser)
    with excel_client() as client:
        response = upload(client, output.getvalue())
    assert response.status_code == 400
    assert "too sparse or large" in response.json()["error"]
    parser.assert_not_called()


def test_text_dates_have_csv_dtype_and_support_time_queries():
    content = workbook_bytes([("Dates", [["event_date", "value"], ["2026-01-01", 3], ["2026-02-02", 4]])])
    frame = parse_xlsx(content)[0][1]
    csv = parse_csv(b"event_date,value\n2026-01-01,3\n2026-02-02,4\n")
    assert str(frame["event_date"].dtype) == str(csv["event_date"].dtype)
    assert pd.api.types.is_datetime64_any_dtype(frame["event_date"])
    llm = Mock()
    llm.generate.return_value = json.dumps({
        "group_by": [{"column": "event_date", "grain": "month"}],
        "aggregations": [{"column": "value", "func": "sum", "alias": "total"}],
    })
    with excel_client(llm) as client:
        dataset = upload(client, content).json()
        response = client.post("/api/query", json={"question": "value per month", "dataset_id": dataset["dataset_id"]})
    assert response.status_code == 200
    assert response.json()["rows"] == [
        {"event_date": "2026-01-01T00:00:00", "total": 3},
        {"event_date": "2026-02-01T00:00:00", "total": 4},
    ]


def test_only_cached_formula_values_are_read_and_network_is_not_used(monkeypatch):
    content = workbook_bytes([("Formulas", [["name", "value"], ["Cached", "=1+2"], ["Uncached", "=WEBSERVICE(\"https://example.com\")"]])])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                xml = ElementTree.fromstring(data)
                namespace = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                xml.find('.//s:c[@r="B2"]/s:v', namespace).text = "3"
                data = ElementTree.tostring(xml)
            target.writestr(entry, data)
    loader = Mock(wraps=load_workbook)
    monkeypatch.setattr("openpyxl.load_workbook", loader)
    no_network = Mock(side_effect=AssertionError("No network access allowed"))
    monkeypatch.setattr("socket.socket.connect", no_network)
    monkeypatch.setattr("urllib.request.urlopen", no_network)
    frame = parse_xlsx(output.getvalue())[0][1]
    assert frame.iloc[0]["value"] == 3
    assert pd.isna(frame.iloc[1]["value"])
    assert loader.call_args.kwargs == ENGINE_OPTIONS
    no_network.assert_not_called()


def test_disguised_macro_workbook_is_rejected():
    content = workbook_bytes([("Sheet", [["value"], [1]])])
    output = BytesIO()
    with ZipFile(BytesIO(content)) as source, ZipFile(output, "w", compression=ZIP_DEFLATED) as target:
        for entry in source.infolist():
            target.writestr(entry, source.read(entry.filename))
        target.writestr("xl/vbaProject.bin", b"macro content")
    with excel_client() as client:
        response = upload(client, output.getvalue())
    assert response.status_code == 400
    assert response.json() == {"error": UNSUPPORTED_EXCEL}


def test_xml_entity_declarations_are_rejected_before_parsing():
    output = BytesIO()
    with ZipFile(output, "w", compression=ZIP_DEFLATED) as archive:
        archive.writestr("xl/workbook.xml", '<!DOCTYPE doc [<!ENTITY x "data">]><doc>&x;</doc>')
    with excel_client() as client:
        response = upload(client, output.getvalue())
    assert response.status_code == 400
    assert response.json() == {"error": "The workbook contains unsupported XML declarations."}


def test_native_excel_times_use_the_csv_time_dtype_and_can_be_queried():
    content = workbook_bytes([("Times", [["customer", "event_time"], ["Ada", time(8, 30)], ["Lin", time(9, 45)]])])
    csv = parse_csv(b"customer,event_time\nAda,08:30:00\nLin,09:45:00\n")
    frame = parse_xlsx(content)[0][1]
    assert str(frame["event_time"].dtype) == str(csv["event_time"].dtype)
    assert frame["event_time"].tolist() == csv["event_time"].tolist()
    llm = Mock()
    llm.generate.return_value = json.dumps({
        "filters": [{"column": "event_time", "op": ">=", "value": "08:45:00"}],
        "select": ["customer"],
    })
    with excel_client(llm) as client:
        response = upload(client, content)
        assert response.status_code == 201
        result = client.post("/api/query", json={"question": "Who arrived after 08:45?", "dataset_id": response.json()["dataset_id"]})
        assert result.status_code == 200
        assert result.json()["rows"] == [{"customer": "Lin"}]


def test_native_times_without_time_column_name_are_json_safe_text():
    content = workbook_bytes([("Shifts", [["customer", "shift"], ["Ada", time(8, 30)], ["Lin", time(9, 45)]])])
    with excel_client() as client:
        response = upload(client, content)
    assert response.status_code == 201
    assert response.json()["rows"] == [{"customer": "Ada", "shift": "08:30:00"}, {"customer": "Lin", "shift": "09:45:00"}]


def test_excel_durations_and_mixed_types_are_json_safe():
    content = workbook_bytes([("Types", [
        ["customer", "duration", "active", "spend", "note"],
        ["Ada", timedelta(hours=27, minutes=5), True, 12.5, None],
        ["Lin", None, False, 20, "NA"],
    ])])
    with excel_client() as client:
        response = upload(client, content)
    assert response.status_code == 201
    assert response.json()["rows"] == [
        {"customer": "Ada", "duration": "P1DT3H5M0S", "active": True, "spend": 12.5, "note": None},
        {"customer": "Lin", "duration": None, "active": False, "spend": 20.0, "note": "NA"},
    ]
