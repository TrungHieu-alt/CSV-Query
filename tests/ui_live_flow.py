"""Test the configured 8080 frontend / 8000 API with real Gemini calls.

Requires the existing frontend to be running. Starts a temporary API if none
is running, without changing .env or credentials. Uses synthetic upload data.
Run: .venv/Scripts/python.exe -B tests/ui_live_flow.py
"""
from __future__ import annotations

from datetime import datetime
from io import BytesIO
import json
from pathlib import Path
import sys
import tempfile
import threading
import time

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from playwright.sync_api import expect, sync_playwright
import requests
import uvicorn

from backend.app.config import Settings
from backend.app.main import create_app


def main():
    artifacts = Path(tempfile.mkdtemp(prefix="csv-live-flow-"))
    report = {"checks": [], "provider": None, "queries": []}
    server = api_thread = None
    try:
        try:
            response = requests.get("http://localhost:8000/api/health", timeout=3)
            assert response.status_code == 200
        except requests.ConnectionError:
            server = uvicorn.Server(uvicorn.Config(create_app(Settings()), host="127.0.0.1", port=8000, log_level="error"))
            api_thread = threading.Thread(target=server.run, daemon=True)
            api_thread.start()
            for _ in range(100):
                if server.started:
                    break
                time.sleep(.05)
            assert server.started, "Temporary API did not start"
            report["checks"].append("Started temporary API on localhost:8000")
        response = requests.get("http://localhost:8080/app.js", timeout=5)
        response.raise_for_status()
        report["checks"].append("Existing frontend serves Excel upload support")
        assert "async function uploadExcel" in response.content.decode("utf-8")
        report["frontend_matches_workspace"] = response.content.decode("utf-8").replace("\r\n", "\n") == (ROOT / "frontend/app.js").read_text(encoding="utf-8").replace("\r\n", "\n")

        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(channel="chrome", headless=True)
            page = browser.new_page(viewport={"width": 1366, "height": 768})
            exceptions = []
            page.on("pageerror", lambda error: exceptions.append(str(error)))
            # No config.js interception: this verifies actual frontend URL,
            # browser CORS, and the configured API and Gemini credentials.
            page.goto("http://localhost:8080")
            expect(page.locator("#row-search")).to_be_enabled(timeout=15000)
            report["checks"].append("Actual 8080 -> 8000 wiring, CORS, schema and rows")
            print("PASS: Actual frontend and backend wiring", flush=True)
            print("Checking configured Gemini connection...", flush=True)
            health = page.request.get("http://localhost:8000/api/health/gemini", timeout=90000)
            report["provider"] = {"status": health.status, "body": health.json()}
            print("Gemini connection:", json.dumps(report["provider"]), flush=True)

            csv_bytes = b"customer,joined_date,spend\nAda,2026-01-02,12.5\nLin,2026-02-03,20\n"
            book = Workbook()
            book.active.title = "Customers"
            book.active.append(["customer", "joined_date", "spend"])
            book.active.append(["Ada", datetime(2026, 1, 2), 12.5])
            book.active.append(["Lin", datetime(2026, 2, 3), 20.0])
            output = BytesIO()
            book.save(output)
            for filename, mime, content in [
                ("live-customers.csv", "text/csv", csv_bytes),
                ("live-customers.xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", output.getvalue()),
            ]:
                page.locator("#file-input").set_input_files({"name": filename, "mimeType": mime, "buffer": content})
                expect(page.locator("#status")).to_contain_text("is ready.", timeout=15000)
                expect(page.locator("#data-summary")).to_have_text("2 rows · 3 columns")
                report["checks"].append(f"Uploaded {filename} via actual app")
                page.locator("#question-input").fill("What is the total spend? Sum the spend column across all rows.")
                with page.expect_response(lambda response: response.url.endswith("/api/query"), timeout=90000) as query_response:
                    page.get_by_role("button", name="Send", exact=True).click()
                response = query_response.value
                result = response.json()
                expect(page.locator("#send-button")).to_be_enabled()
                query = {"filename": filename, "status": response.status}
                if response.status == 200:
                    assert "rows" in result and "plan" in result, "Gemini did not produce an executable query"
                    values = [value for row in result["rows"] for value in row.values() if isinstance(value, (int, float))]
                    assert values == [32.5], "Live query produced an incorrect sum"
                    expect(page.locator(".message.assistant").last).to_contain_text("$32.50")
                    query["result"] = result
                    print("PASS: Real Gemini question on", filename, flush=True)
                else:
                    query["error"] = result
                    expect(page.locator(".message.assistant.error").last).to_be_visible()
                    print("BLOCKED: Real Gemini question on", filename, response.status, result.get("error_code"), flush=True)
                report["queries"].append(query)
                page.screenshot(path=str(artifacts / (filename + ".png")))

            if all(query["status"] == 200 for query in report["queries"]):
                page.locator("#question-input").fill("Sum spend for each month of joined_date. Group by month.")
                with page.expect_response(lambda response: response.url.endswith("/api/query"), timeout=90000) as monthly_response:
                    page.get_by_role("button", name="Send", exact=True).click()
                response = monthly_response.value
                assert response.status == 200, "Live monthly query failed"
                result = response.json()
                values = sorted(value for row in result["rows"] for value in row.values() if isinstance(value, (int, float)))
                assert values == [12.5, 20.0]
                assert result["result_type"] == "time_series"
                expect(page.locator(".message.assistant").last.locator("canvas")).to_be_visible()
                report["queries"].append({"filename": "live-customers.xlsx", "question": "monthly spend", "status": response.status, "result": result})
                print("PASS: Real Gemini monthly question renders Excel line chart", flush=True)

                multi = Workbook()
                multi.remove(multi.active)
                for index, name in enumerate(["First", "Second", "Third"], 1):
                    sheet = multi.create_sheet(name)
                    sheet.append(["customer", "spend"])
                    sheet.append([f"Customer {index}", index * 10])
                multi_output = BytesIO()
                multi.save(multi_output)
                page.locator("#file-input").set_input_files({"name": "live-three.xlsx", "mimeType": "application/octet-stream", "buffer": multi_output.getvalue()})
                expect(page.locator("#status")).to_have_text('“live-three.xlsx – First” is ready.')
                page.locator("#dataset-select").select_option(label="live-three.xlsx – Third")
                expect(page.locator("#row-search")).to_be_enabled()
                page.locator("#question-input").fill("What is the total spend? Sum the spend column across all rows.")
                with page.expect_response(lambda response: response.url.endswith("/api/query"), timeout=90000) as third_response:
                    page.get_by_role("button", name="Send", exact=True).click()
                response = third_response.value
                assert response.status == 200, "Live question after switching sheets failed"
                result = response.json()
                values = [value for row in result["rows"] for value in row.values() if isinstance(value, (int, float))]
                assert values == [30]
                expect(page.locator(".message.assistant").last).to_contain_text("$30.00")
                report["queries"].append({"filename": "live-three.xlsx", "question": "third sheet total", "status": response.status, "result": result})
                page.screenshot(path=str(artifacts / "live-three.xlsx.png"))
                print("PASS: Real Gemini question after switching to third worksheet", flush=True)
            assert not exceptions, exceptions
            report["checks"].append("No browser JavaScript exceptions")
            browser.close()
        report["live_queries_passed"] = all(query["status"] == 200 for query in report["queries"])
        print("Report:", artifacts / "results.json", flush=True)
        print("Live queries passed:", report["live_queries_passed"], flush=True)
    finally:
        (artifacts / "results.json").write_text(json.dumps(report, indent=2), encoding="utf-8")
        if server:
            server.should_exit = True
            api_thread.join(timeout=5)


if __name__ == "__main__":
    main()
